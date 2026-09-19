"""Create a test draft .docx document (unformatted, to be optimized)."""

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Add a title-like paragraph (will be mapped to Title style)
doc.add_paragraph("论人工智能在医疗领域的应用研究")

# Heading 1
doc.add_heading("引言", level=1)

# Body paragraphs (plain text, no special formatting)
body_texts = [
    "人工智能技术在医疗领域的应用越来越广泛，我觉得这个趋势还会继续发展下去。",
    "目前主要的AI医疗应用包括医学影像识别、辅助诊断、药物研发等几个方面。",
    "由于技术还不太成熟的原因，所以在实际应用中还面临很多挑战。",
    "本文将从技术现状、应用场景和未来展望三个方面进行论述。",
]

for text in body_texts:
    doc.add_paragraph(text)

# Heading 2
doc.add_heading("技术现状", level=2)

more_body = [
    "深度学习是当前AI医疗的核心技术，即在大规模数据训练下能获得很好的效果。",
    "然而数据隐私和模型可解释性问题仍是有待解决的关键难点。",
]
for text in more_body:
    doc.add_paragraph(text)

# A simple table
table = doc.add_table(rows=3, cols=2)
table.style = "Table Grid"
data = [
    ["技术名称", "应用领域"],
    ["CNN", "医学影像"],
    ["NLP", "病历分析"],
]
for i, row_data in enumerate(data):
    for j, cell_text in enumerate(row_data):
        table.rows[i].cells[j].text = cell_text

output = r"C:\Users\lenpvp\AppData\Roaming\TRAE SOLO CN\ModularData\ai-agent\work-mode-projects\6a891b8548fd7c69dff220af\doc-format-optimizer\backend\test_draft.docx"
doc.save(output)
print(f"Test draft created: {output}")

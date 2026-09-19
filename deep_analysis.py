"""Deep analysis: extract all text, red text, text boxes, and format hints."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt, Emu
from lxml import etree

docx_path = "user_uploaded.docx"
doc = Document(docx_path)

print("=" * 70)
print("  深度分析：.doc 转换后的 .docx 文档")
print("=" * 70)

# 1. Extract ALL paragraphs with color info
print("\n【1】所有带颜色的段落:")
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if not text:
        continue
    
    for run in para.runs:
        color = None
        # Check run color via XML
        rpr = run._element.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr')
        if rpr is not None:
            color_elem = rpr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}color')
            if color_elem is not None:
                color = color_elem.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
        
        if color and color != 'auto' and color != '000000':
            print(f"  [Para {i}] color=#{color} font={run.font.name} size={run.font.size} | {text[:80]}")
            break

# 2. Extract text boxes / shapes (wps:txbx or mc:AlternateContent)
print("\n【2】文本框/形状内容:")
nsmap = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'wps': 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape',
    'mc': 'http://schemas.openxmlformats.org/markup-compatibility/2006',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'v': 'urn:schemas-microsoft-com:vml',
    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
}

# Search for text boxes in the entire document XML
body = doc.element.body
# Find all w:txbxContent elements
txbx_ns = 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape'
alt_ns = 'http://schemas.openxmlformats.org/markup-compatibility/2006'

text_boxes = body.findall('.//{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}txbxContent')
if not text_boxes:
    # Try VML textboxes
    text_boxes = body.findall('.//{urn:schemas-microsoft-com:vml}textbox')

print(f"  找到 {len(text_boxes)} 个文本框")
for idx, tb in enumerate(text_boxes):
    # Extract text from the textbox
    texts = tb.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')
    full_text = ''.join(t.text or '' for t in texts)
    if full_text.strip():
        print(f"  文本框[{idx}]: {full_text[:120]}")

# 3. Search for format requirement keywords in all text
print("\n【3】包含格式要求关键词的段落:")
keywords = ['字体', '字号', '行距', '页边距', '缩进', '对齐', '加粗', '宋体', '黑体', 
            '楷体', '三号', '四号', '五号', '小四', '小五', '二号', 'A4', '页眉', '页脚',
            '磅', '厘米', '倍行距', '首行', '段前', '段后', '目录', '摘要', '关键词',
            '参考文献', '正文', '标题', '图表']

for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if not text:
        continue
    matched = [kw for kw in keywords if kw in text]
    if matched:
        style = para.style.name if para.style else "?"
        print(f"  [Para {i}] style={style} matched={matched}")
        print(f"    text: {text[:100]}")

# 4. Extract table contents
print("\n【4】表格内容:")
for tidx, table in enumerate(doc.tables):
    print(f"  表格[{tidx}]: {len(table.rows)}行 x {len(table.columns)}列")
    for ridx, row in enumerate(table.rows):
        cells = [cell.text.strip()[:30] for cell in row.cells]
        print(f"    行{ridx}: {cells}")
        if ridx >= 5:
            print(f"    ... (共{len(table.rows)}行)")
            break

# 5. Section properties
print("\n【5】节属性:")
sections = doc.sections
for sidx, section in enumerate(sections):
    print(f"  节[{sidx}]:")
    print(f"    页面: {section.page_width} x {section.page_height} EMU")
    if section.page_width:
        print(f"    页面(pt): {section.page_width / 12700:.1f} x {section.page_height / 12700:.1f}")
    print(f"    边距(pt): 上={section.top_margin / 12700 if section.top_margin else '?'}, 下={section.bottom_margin / 12700 if section.bottom_margin else '?'}, 左={section.left_margin / 12700 if section.left_margin else '?'}, 右={section.right_margin / 12700 if section.right_margin else '?'}")
    # Header
    header = section.header
    if header and header.paragraphs:
        header_text = ' | '.join(p.text for p in header.paragraphs if p.text.strip())
        if header_text:
            print(f"    页眉: {header_text[:80]}")
    footer = section.footer
    if footer and footer.paragraphs:
        footer_text = ' | '.join(p.text for p in footer.paragraphs if p.text.strip())
        if footer_text:
            print(f"    页脚: {footer_text[:80]}")

# 6. Full text dump (first 60 paragraphs)
print("\n【6】前60段完整文本:")
for i, para in enumerate(doc.paragraphs[:60]):
    text = para.text.strip()
    if text:
        print(f"  [{i}] ({para.style.name}) {text[:100]}")

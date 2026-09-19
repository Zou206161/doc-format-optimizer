"""Create a test .docx template that simulates a thesis format."""

from docx import Document
from docx.shared import Pt, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

doc = Document()

# Page setup: A4
section = doc.sections[0]
section.page_width = Emu(int(595.3 * 12700))
section.page_height = Emu(int(841.9 * 12700))
section.top_margin = Pt(72)
section.bottom_margin = Pt(72)
section.left_margin = Pt(90)
section.right_margin = Pt(90)

# Header
header = section.header
header.paragraphs[0].text = "本科毕业论文"
header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

# Title (黑体 二号 居中)
title = doc.add_paragraph()
title.style = doc.styles['Title']
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title.add_run("基于深度学习的图像识别研究")
title_run.font.size = Pt(22)
title_run.font.name = "SimHei"
rPr = title_run._element.get_or_add_rPr()
rFonts = rPr.makeelement(qn("w:rFonts"), {})
rFonts.set(qn("w:eastAsia"), "黑体")
rFonts.set(qn("w:ascii"), "SimHei")
rFonts.set(qn("w:hAnsi"), "SimHei")
rPr.insert(0, rFonts)
title_run.font.bold = True

# Heading 1 (黑体 三号 居中 加粗)
h1 = doc.add_heading("第一章 绪论", level=1)
h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in h1.runs:
    run.font.size = Pt(16)
    run.font.name = "SimHei"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), "黑体")
    rFonts.set(qn("w:ascii"), "SimHei")
    rFonts.set(qn("w:hAnsi"), "SimHei")
    run.font.bold = True
h1.paragraph_format.space_before = Pt(24)
h1.paragraph_format.space_after = Pt(18)

# Heading 2 (黑体 四号 左对齐 加粗)
h2 = doc.add_heading("1.1 研究背景", level=2)
h2.alignment = WD_ALIGN_PARAGRAPH.LEFT
for run in h2.runs:
    run.font.size = Pt(14)
    run.font.name = "SimHei"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), "黑体")
    rFonts.set(qn("w:ascii"), "SimHei")
    rFonts.set(qn("w:hAnsi"), "SimHei")
    run.font.bold = True
h2.paragraph_format.space_before = Pt(18)
h2.paragraph_format.space_after = Pt(12)

# Body text (宋体 小四 1.5倍行距 首行缩进2字符 两端对齐)
body_texts = [
    "随着人工智能技术的快速发展，深度学习在计算机视觉领域取得了显著的成果。图像识别作为其中的重要分支，已广泛应用于自动驾驶、医学影像分析等场景。",
    "本文旨在研究基于卷积神经网络的图像识别方法，探讨不同网络结构对识别精度的影响，并提出一种改进的注意力机制模块以提升模型性能。",
    "实验结果表明，所提方法在标准数据集上的准确率达到了95.3%，相比基线模型提升了3.2个百分点，验证了该方法的有效性。",
]

for text in body_texts:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = para.add_run(text)
    run.font.size = Pt(12)
    run.font.name = "Times New Roman"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), "宋体")
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")

    pf = para.paragraph_format
    pf.line_spacing = 1.5
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.first_line_indent = Pt(24)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

# Heading 2
h2b = doc.add_heading("1.2 研究意义", level=2)
h2b.alignment = WD_ALIGN_PARAGRAPH.LEFT
for run in h2b.runs:
    run.font.size = Pt(14)
    run.font.name = "SimHei"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), "黑体")
    rFonts.set(qn("w:ascii"), "SimHei")
    rFonts.set(qn("w:hAnsi"), "SimHei")
    run.font.bold = True

body2 = doc.add_paragraph()
body2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
run2 = body2.add_run("本研究的意义在于为图像识别领域提供了新的思路和方法，具有一定的理论价值和实际应用前景。")
run2.font.size = Pt(12)
run2.font.name = "Times New Roman"
rPr = run2._element.get_or_add_rPr()
rFonts = rPr.find(qn("w:rFonts"))
if rFonts is None:
    rFonts = rPr.makeelement(qn("w:rFonts"), {})
    rPr.insert(0, rFonts)
rFonts.set(qn("w:eastAsia"), "宋体")
rFonts.set(qn("w:ascii"), "Times New Roman")
rFonts.set(qn("w:hAnsi"), "Times New Roman")
body2.paragraph_format.line_spacing = 1.5
body2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
body2.paragraph_format.first_line_indent = Pt(24)

# A table
table = doc.add_table(rows=2, cols=3)
table.style = "Table Grid"
for i, row_data in enumerate([["方法", "准确率", "参数量"], ["ResNet-50", "94.2%", "25.6M"]]):
    for j, cell_text in enumerate(row_data):
        table.rows[i].cells[j].text = cell_text

output = r"C:\Users\lenpvp\AppData\Roaming\TRAE SOLO CN\ModularData\ai-agent\work-mode-projects\6a891b8548fd7c69dff220af\doc-format-optimizer\backend\test_template.docx"
doc.save(output)
print(f"Test template created: {output}")

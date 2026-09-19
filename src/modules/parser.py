"""Document parser — reads a .docx file and builds a DocumentObject."""

from __future__ import annotations

import logging

from docx import Document
from docx.shared import Emu, Pt
from lxml import etree

from src.models import (
    Alignment,
    DocumentObject,
    FontFormat,
    ParagraphFormat,
    ParagraphInfo,
    TableInfo,
)

logger = logging.getLogger(__name__)

_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

_HEADING_STYLE_MAP = {
    "heading 1": 1, "heading 2": 2, "heading 3": 3, "heading 4": 4,
    "标题 1": 1, "标题 2": 2, "标题 3": 3, "标题 4": 4,
    "title": 1, "标题": 1,
    "subtitle": 1, "副标题": 1,
}


def _emu_to_pt(value) -> float | None:
    if value is None:
        return None
    try:
        return round(Emu(value).pt, 2)
    except Exception:
        return None


def _extract_font_from_run(run) -> FontFormat:
    font = run.font
    fmt = FontFormat(
        bold=font.bold,
        italic=font.italic,
        underline=font.underline,
    )
    if font.name:
        fmt.font_name = font.name
    if font.size:
        fmt.font_size = round(font.size.pt, 2)
    if font.color and font.color.rgb:
        fmt.color = str(font.color.rgb)

    rpr = run._element.find(".//w:rPr", _NS)
    if rpr is not None:
        rfonts = rpr.find("w:rFonts", _NS)
        if rfonts is not None:
            east = rfonts.get(f"{{{_NS['w']}}}eastAsia")
            if east:
                fmt.font_name_east_asia = east
            if not fmt.font_name:
                ascii_font = rfonts.get(f"{{{_NS['w']}}}ascii")
                if ascii_font:
                    fmt.font_name = ascii_font
    return fmt


def _merge_run_fonts(runs) -> FontFormat:
    """Merge font info from multiple runs — first non-None value wins for each attribute."""
    if not runs:
        return FontFormat()
    merged = FontFormat()
    for run in runs:
        rf = _extract_font_from_run(run)
        for attr in ("font_name", "font_name_east_asia", "font_size", "bold", "italic", "underline", "color"):
            val = getattr(rf, attr)
            if val is not None and getattr(merged, attr) is None:
                setattr(merged, attr, val)
    return merged


_KNOWN_HEADINGS = {
    "引言", "绪论", "前言", "导言", "结论", "结语", "总结", "参考文献",
    "参考资料", "致谢", "附录", "摘要", "关键词", "abstract", "introduction",
    "conclusion", "references", "acknowledgments",
}


def _is_short_non_sentence(text: str, max_len: int = 20) -> bool:
    """Check if text is short and doesn't end with sentence punctuation."""
    stripped = text.strip()
    if not stripped or len(stripped) > max_len:
        return False
    return not stripped.endswith((
        "。", "！", "？", ".", "!", "?", ";", "；", "，", ",",
        "）", "」", "』", "】", ")", "}", "]",
    ))


def _detect_heading(style_name: str, font_fmt: FontFormat | None = None, text: str = "") -> tuple[bool, int]:
    """Detect heading by style name; fall back to heuristic if no style match."""
    name = style_name.strip().lower()
    if name in _HEADING_STYLE_MAP:
        return True, _HEADING_STYLE_MAP[name]

    stripped = text.strip()

    # Heuristic 1: bold + larger font + short text without sentence punctuation
    if font_fmt and stripped:
        if (font_fmt.bold and font_fmt.font_size and font_fmt.font_size >= 14
                and len(stripped) <= 40
                and not stripped.endswith(("。", "！", "？", ".", "!", "?", ";", "；", "，", ","))):
            if font_fmt.font_size >= 18:
                return True, 1
            elif font_fmt.font_size >= 15:
                return True, 2
            else:
                return True, 3

    # Heuristic 2: known heading keywords
    if stripped in _KNOWN_HEADINGS:
        return True, 1

    # Heuristic 3: short non-sentence text when document has no formatting differentiation
    # (font_fmt has no bold/size info, meaning all paragraphs inherit from Normal style)
    if stripped and _is_short_non_sentence(stripped, max_len=15):
        no_fmt = (font_fmt is None or
                  (font_fmt.bold is None and font_fmt.font_size is None))
        if no_fmt:
            return True, 1

    return False, 0


def _extract_paragraph_format(para) -> ParagraphFormat:
    pf = para.paragraph_format
    fmt = ParagraphFormat()

    if pf.alignment is not None:
        align_map = {
            0: Alignment.LEFT,
            1: Alignment.CENTER,
            2: Alignment.RIGHT,
            3: Alignment.JUSTIFY,
        }
        fmt.alignment = align_map.get(int(pf.alignment))

    if pf.line_spacing is not None:
        rule_str = str(pf.line_spacing_rule).upper() if pf.line_spacing_rule is not None else ""
        if "EXACT" in rule_str:
            fmt.line_spacing = _emu_to_pt(int(pf.line_spacing))
            fmt.line_spacing_rule = "exact"
        elif "AT_LEAST" in rule_str:
            fmt.line_spacing = _emu_to_pt(int(pf.line_spacing))
            fmt.line_spacing_rule = "at_least"
        else:
            fmt.line_spacing = round(float(pf.line_spacing), 2)
            fmt.line_spacing_rule = "multiple"

    fmt.space_before = _emu_to_pt(pf.space_before)
    fmt.space_after = _emu_to_pt(pf.space_after)
    fmt.first_line_indent = _emu_to_pt(pf.first_line_indent)
    fmt.left_indent = _emu_to_pt(pf.left_indent)
    fmt.right_indent = _emu_to_pt(pf.right_indent)
    return fmt


def _extract_header_footer(section) -> tuple[str | None, str | None]:
    header_text = None
    footer_text = None
    try:
        hdr = section.header
        if hdr and hdr.paragraphs:
            parts = [p.text for p in hdr.paragraphs if p.text.strip()]
            if parts:
                header_text = "\n".join(parts)
    except Exception:
        pass
    try:
        ftr = section.footer
        if ftr and ftr.paragraphs:
            parts = [p.text for p in ftr.paragraphs if p.text.strip()]
            if parts:
                footer_text = "\n".join(parts)
    except Exception:
        pass
    return header_text, footer_text


def parse_document(file_path: str) -> DocumentObject:
    """Open a .docx file and return a fully populated DocumentObject."""
    doc = Document(file_path)

    paragraphs: list[ParagraphInfo] = []
    for idx, para in enumerate(doc.paragraphs):
        style_name = para.style.name if para.style else "Normal"
        para_fmt = _extract_paragraph_format(para)

        runs = para.runs
        font_fmt = _merge_run_fonts(runs) if runs else FontFormat()

        is_heading, level = _detect_heading(style_name, font_fmt, para.text)

        paragraphs.append(
            ParagraphInfo(
                index=idx,
                text=para.text,
                style_name=style_name,
                is_heading=is_heading,
                heading_level=level,
                font=font_fmt,
                paragraph=para_fmt,
            )
        )

    tables: list[TableInfo] = []
    for tbl in doc.tables:
        rows = len(tbl.rows)
        cols = len(tbl.columns) if tbl.rows else 0
        cells: list[list[str]] = []
        for row in tbl.rows:
            cells.append([cell.text for cell in row.cells])
        tables.append(TableInfo(rows=rows, cols=cols, cells=cells))

    doc_obj = DocumentObject(
        paragraphs=paragraphs,
        tables=tables,
    )

    if doc.sections:
        section = doc.sections[0]
        if section.page_width:
            doc_obj.page_width = _emu_to_pt(section.page_width)
        if section.page_height:
            doc_obj.page_height = _emu_to_pt(section.page_height)
        doc_obj.margin_top = _emu_to_pt(section.top_margin)
        doc_obj.margin_bottom = _emu_to_pt(section.bottom_margin)
        doc_obj.margin_left = _emu_to_pt(section.left_margin)
        doc_obj.margin_right = _emu_to_pt(section.right_margin)
        doc_obj.header, doc_obj.footer = _extract_header_footer(section)

    return doc_obj

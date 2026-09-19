"""用用户的真实文档测试格式优化全流程。"""
import os, sys, json
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from src.modules.parser import parse_document
from src.modules.format_applier import apply_format
from src.modules.doc_generator import generate_document
from src.modules.requirement_parser import parse_requirement
from docx import Document
from docx.shared import Emu

BASE = os.path.dirname(os.path.abspath(__file__))
USER_DOC = os.path.join(BASE, "user_real_doc.docx")

USER_REQUIREMENTS = """学术论文格式规范要求如下：
纸张规格使用A4纸张。
页边距要求上边距二点五厘米，下边距二点五厘米，左边距三厘米，右边距二点五厘米。
中文字体正文使用宋体，大标题使用黑体。
英文字体和数字使用Times New Roman字体。
字号标准规定大标题使用二号字，一级标题使用三号字，二级标题使用四号字，三级标题使用小四号字，正文内容使用小四号字，图表标题使用五号字。
段落格式要求正文段落首行缩进两个汉字字符。
全文行距设定为固定值二十磅，段前段后间距均调整为零。
标题位置要求论文大标题居中对齐，一级标题顶格，二级标题缩进两个字符，三级标题缩进四个字符。
图表规范要求图表需居中放置，图名置于图下方，表名置于表上方，均使用五号字，图表编号采用阿拉伯数字连续编号，例如图一、表一。
参考文献著录采用GB/T7714格式，正文中的引用处以数字上标标注序号，例如一，文末参考文献列表按照引用顺序排列。"""

print("=" * 70)
print("  用户真实文档测试")
print("=" * 70)

PASS = 0
FAIL = 0

def check(name, condition, detail=""):
    global PASS, FAIL
    if condition:
        print(f"  ✅ {name}")
        PASS += 1
    else:
        print(f"  ❌ {name} — {detail}")
        FAIL += 1

# ============================================================
# 1. 解析文档
# ============================================================
print("\n【1. 文档解析】")
doc_obj = parse_document(USER_DOC)
print(f"  段落数: {len(doc_obj.paragraphs)}")
print(f"  表格数: {len(doc_obj.tables)}")
print(f"  页面: {doc_obj.page_width}x{doc_obj.page_height}pt")

headings = [(i, p.text, p.is_heading, p.heading_level) for i, p in enumerate(doc_obj.paragraphs) if p.is_heading]
print(f"\n  检测到 {len(headings)} 个标题:")
for idx, text, is_h, level in headings:
    print(f"    [{idx}] L{level}: {text[:40]}")

# 检查关键标题是否被检测到
title_texts = [h[1] for h in headings if h[3] == 0 or True]  # all headings
check("检测到标题", len(headings) > 0, f"仅检测到{len(headings)}个")
check("检测到'引言'", any("引言" in h[1] for h in headings))
check("检测到'结论'", any("结论" in h[1] for h in headings))
check("检测到'参考文献'", any("参考文献" in h[1] for h in headings))
check("检测到'责任归属'", any("责任归属" in h[1] for h in headings))
check("检测到'数据隐私'", any("数据隐私" in h[1] for h in headings))
check("检测到'算法公平'", any("算法公平" in h[1] for h in headings))

# ============================================================
# 2. 文字要求解析
# ============================================================
print("\n【2. 文字要求解析（模式匹配）】")
rule_set = parse_requirement(USER_REQUIREMENTS)

print(f"  样式: {list(rule_set.styles.keys())}")
dl = rule_set.document_level
print(f"\n  页面设置:")
print(f"    纸张: {dl.page_width}x{dl.page_height}pt")
print(f"    上边距: {dl.margin_top}pt (要求70.9)")
print(f"    下边距: {dl.margin_bottom}pt (要求70.9)")
print(f"    左边距: {dl.margin_left}pt (要求85.0)")
print(f"    右边距: {dl.margin_right}pt (要求70.9)")

check("A4纸张", dl.page_width == 595.3 and dl.page_height == 841.9)
check("上边距=2.5cm", abs(dl.margin_top - 70.88) < 1, f"实际={dl.margin_top}")
check("下边距=2.5cm", abs(dl.margin_bottom - 70.88) < 1, f"实际={dl.margin_bottom}")
check("左边距=3cm", abs(dl.margin_left - 85.05) < 1, f"实际={dl.margin_left}")
check("右边距=2.5cm", abs(dl.margin_right - 70.88) < 1, f"实际={dl.margin_right}")

print(f"\n  Title 样式:")
title_s = rule_set.styles.get("Title")
if title_s:
    print(f"    字体: {title_s.font.font_name_east_asia} (要求黑体)")
    print(f"    字号: {title_s.font.font_size}pt (要求22pt)")
    print(f"    加粗: {title_s.font.bold} (要求True)")
    print(f"    对齐: {title_s.paragraph.alignment} (要求center)")
    check("Title字体=黑体", title_s.font.font_name_east_asia == "黑体")
    check("Title字号=22pt(二号)", title_s.font.font_size == 22.0)
    check("Title加粗", title_s.font.bold == True)
    check("Title居中", title_s.paragraph.alignment == "center")
else:
    check("Title样式存在", False, "Title样式不存在")

print(f"\n  Heading1 样式:")
h1_s = rule_set.styles.get("Heading1")
if h1_s:
    print(f"    字体: {h1_s.font.font_name_east_asia}")
    print(f"    字号: {h1_s.font.font_size}pt (要求16pt)")
    print(f"    对齐: {h1_s.paragraph.alignment}")
    print(f"    缩进: {h1_s.paragraph.first_line_indent} (要求0=顶格)")
    print(f"    段前: {h1_s.paragraph.space_before} 段后: {h1_s.paragraph.space_after}")
    check("H1字号=16pt(三号)", h1_s.font.font_size == 16.0)
    check("H1顶格(缩进=0)", h1_s.paragraph.first_line_indent == 0.0, f"实际={h1_s.paragraph.first_line_indent}")
    check("H1段前=0", h1_s.paragraph.space_before == 0.0)
    check("H1段后=0", h1_s.paragraph.space_after == 0.0)

print(f"\n  Heading2 样式:")
h2_s = rule_set.styles.get("Heading2")
if h2_s:
    print(f"    字号: {h2_s.font.font_size}pt (要求14pt)")
    print(f"    缩进: {h2_s.paragraph.first_line_indent}pt (要求24pt=2字符)")
    check("H2字号=14pt(四号)", h2_s.font.font_size == 14.0)
    check("H2缩进=24pt(2字符)", h2_s.paragraph.first_line_indent == 24.0, f"实际={h2_s.paragraph.first_line_indent}")

print(f"\n  Heading3 样式:")
h3_s = rule_set.styles.get("Heading3")
if h3_s:
    print(f"    字号: {h3_s.font.font_size}pt (要求12pt)")
    print(f"    缩进: {h3_s.paragraph.first_line_indent}pt (要求48pt=4字符)")
    check("H3字号=12pt(小四)", h3_s.font.font_size == 12.0)
    check("H3缩进=48pt(4字符)", h3_s.paragraph.first_line_indent == 48.0, f"实际={h3_s.paragraph.first_line_indent}")

print(f"\n  Body 样式:")
body_s = rule_set.styles.get("Body")
if body_s:
    print(f"    字体: {body_s.font.font_name_east_asia} / {body_s.font.font_name}")
    print(f"    字号: {body_s.font.font_size}pt (要求12pt)")
    print(f"    行距: {body_s.paragraph.line_spacing} 规则: {body_s.paragraph.line_spacing_rule} (要求20pt exact)")
    print(f"    缩进: {body_s.paragraph.first_line_indent}pt (要求24pt)")
    print(f"    段前: {body_s.paragraph.space_before} 段后: {body_s.paragraph.space_after}")
    check("Body字体=宋体", body_s.font.font_name_east_asia == "宋体")
    check("Body字号=12pt(小四)", body_s.font.font_size == 12.0)
    check("Body行距=20pt固定", body_s.paragraph.line_spacing == 20.0 and body_s.paragraph.line_spacing_rule == "exact",
          f"实际={body_s.paragraph.line_spacing}/{body_s.paragraph.line_spacing_rule}")
    check("Body缩进=24pt", body_s.paragraph.first_line_indent == 24.0)
    check("Body段前=0", body_s.paragraph.space_before == 0.0)
    check("Body段后=0", body_s.paragraph.space_after == 0.0)

# ============================================================
# 3. 生成成品文档
# ============================================================
print("\n【3. 生成成品文档】")
formatted = apply_format(doc_obj, None, rule_set)

print(f"  格式化段落数: {len(formatted.paragraphs)}")
style_counts = {}
for fp in formatted.paragraphs:
    style_counts[fp.style_key] = style_counts.get(fp.style_key, 0) + 1
print(f"  样式分布: {style_counts}")

# 检查标题是否映射为 Title
title_paras = [fp for fp in formatted.paragraphs if fp.style_key == "Title"]
check("有Title段落", len(title_paras) > 0, "无Title段落")
if title_paras:
    print(f"    Title文本: {title_paras[0].text[:50]}")

# 检查Heading1段落
h1_paras = [fp for fp in formatted.paragraphs if fp.style_key == "Heading1"]
check("有Heading1段落", len(h1_paras) > 0, "无Heading1段落")
print(f"    H1段落数: {len(h1_paras)}")
for h in h1_paras:
    print(f"      '{h.text[:30]}'")

out_path = os.path.join(BASE, "user_final_output.docx")
generate_document(formatted, out_path)
print(f"\n  生成: {out_path} ({os.path.getsize(out_path)} bytes)")

# ============================================================
# 4. 验证成品文档
# ============================================================
print("\n【4. 成品文档验证】")
v = parse_document(out_path)
print(f"  段落数: {len(v.paragraphs)}")
print(f"  页面: {v.page_width}x{v.page_height}pt")
print(f"  边距: 上{v.margin_top} 下{v.margin_bottom} 左{v.margin_left} 右{v.margin_right}")

check("成品页面=A4", v.page_width == 595.3 and v.page_height == 841.9)
check("成品上边距≈70.9", v.margin_top is not None and abs(v.margin_top - 70.88) < 1, f"实际={v.margin_top}")
check("成品下边距≈70.9", v.margin_bottom is not None and abs(v.margin_bottom - 70.88) < 1, f"实际={v.margin_bottom}")
check("成品左边距≈85", v.margin_left is not None and abs(v.margin_left - 85.05) < 1, f"实际={v.margin_left}")
check("成品右边距≈70.9", v.margin_right is not None and abs(v.margin_right - 70.88) < 1, f"实际={v.margin_right}")

# 找标题段落
title_para = None
for p in v.paragraphs:
    if "人工智能在医疗影像诊断" in p.text:
        title_para = p
        break

if title_para:
    print(f"\n  标题段落: '{title_para.text[:40]}'")
    print(f"    字体: {title_para.font.font_name_east_asia} (要求黑体)")
    print(f"    字号: {title_para.font.font_size}pt (要求22pt)")
    print(f"    加粗: {title_para.font.bold} (要求True)")
    align_str = str(title_para.paragraph.alignment) if title_para.paragraph.alignment else "None"
    print(f"    对齐: {align_str} (要求center)")
    check("成品标题字体=黑体", title_para.font.font_name_east_asia == "黑体", f"实际={title_para.font.font_name_east_asia}")
    check("成品标题字号=22pt", title_para.font.font_size == 22.0, f"实际={title_para.font.font_size}")
    check("成品标题加粗", title_para.font.bold == True, f"实际={title_para.font.bold}")
    check("成品标题居中", "center" in align_str.lower(), f"实际={align_str}")

# 找正文段落
body_para = None
for p in v.paragraphs:
    if "随着深度学习" in p.text or "人工智能技术在医疗影像" in p.text:
        body_para = p
        break

if body_para:
    print(f"\n  正文段落: '{body_para.text[:40]}'")
    print(f"    字体: {body_para.font.font_name_east_asia} (要求宋体)")
    print(f"    字号: {body_para.font.font_size}pt (要求12pt)")
    align_str = str(body_para.paragraph.alignment) if body_para.paragraph.alignment else "None"
    print(f"    对齐: {align_str} (要求justify)")
    print(f"    行距: {body_para.paragraph.line_spacing} 规则: {body_para.paragraph.line_spacing_rule}")
    print(f"    缩进: {body_para.paragraph.first_line_indent}pt (要求24pt)")
    check("成品正文字体=宋体", body_para.font.font_name_east_asia == "宋体", f"实际={body_para.font.font_name_east_asia}")
    check("成品正文字号=12pt", body_para.font.font_size == 12.0, f"实际={body_para.font.font_size}")
    check("成品正文行距=20pt固定", body_para.paragraph.line_spacing == 20.0 and "exact" in str(body_para.paragraph.line_spacing_rule).lower(),
          f"实际={body_para.paragraph.line_spacing}/{body_para.paragraph.line_spacing_rule}")
    check("成品正文缩进=24pt", body_para.paragraph.first_line_indent == 24.0, f"实际={body_para.paragraph.first_line_indent}")

# 找一级标题段落
h1_para = None
for p in v.paragraphs:
    if p.text.strip() == "引言":
        h1_para = p
        break

if h1_para:
    print(f"\n  一级标题: '{h1_para.text}'")
    print(f"    字体: {h1_para.font.font_name_east_asia} (要求黑体)")
    print(f"    字号: {h1_para.font.font_size}pt (要求16pt)")
    print(f"    缩进: {h1_para.paragraph.first_line_indent}pt (要求0=顶格)")
    check("成品H1字体=黑体", h1_para.font.font_name_east_asia == "黑体", f"实际={h1_para.font.font_name_east_asia}")
    check("成品H1字号=16pt", h1_para.font.font_size == 16.0, f"实际={h1_para.font.font_size}")
    check("成品H1顶格", h1_para.paragraph.first_line_indent == 0.0 or h1_para.paragraph.first_line_indent is None,
          f"实际={h1_para.paragraph.first_line_indent}")

print(f"\n{'=' * 70}")
print(f"  测试结果: 通过={PASS} 失败={FAIL} 总计={PASS+FAIL}")
print(f"{'=' * 70}")
